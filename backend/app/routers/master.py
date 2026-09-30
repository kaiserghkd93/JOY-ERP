from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from app.database import get_db
from app.models.master import Item, Partner
from app.schemas.master import ItemCreate, ItemUpdate, ItemOut, PartnerCreate, PartnerUpdate, PartnerOut
import openpyxl, io

router = APIRouter(prefix="/master", tags=["기준정보"])


@router.post("/items", response_model=ItemOut)
def create_item(body: ItemCreate, db: Session = Depends(get_db)):
    if db.get(Item, body.part_no):
        raise HTTPException(409, f"이미 존재: {body.part_no}")
    item = Item(**body.model_dump())
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


@router.get("/items/customer-parts")
def customer_parts(db: Session = Depends(get_db)):
    """고객사별 품번 목록 반환 {partner_id: [part_no, ...]}"""
    from app.models.sales import SalesOrder
    rows = db.query(SalesOrder.partner_id, SalesOrder.part_no).distinct().all()
    result: dict[str, list[str]] = {}
    for partner_id, part_no in rows:
        result.setdefault(partner_id, []).append(part_no)
    return result


@router.get("/items", response_model=list[ItemOut])
def list_items(active_only: bool = True, partner_id: str | None = None, db: Session = Depends(get_db)):
    q = db.query(Item)
    if active_only:
        q = q.filter(Item.active == True)
    if partner_id:
        from app.models.sales import SalesOrder, Shipment
        from app.models.purchase import PurchaseOrder
        so_parts = {r[0] for r in db.query(SalesOrder.part_no).filter(SalesOrder.partner_id == partner_id).all()}
        sh_parts = {r[0] for r in db.query(Shipment.part_no).join(SalesOrder, SalesOrder.so_no == Shipment.so_no).filter(SalesOrder.partner_id == partner_id).all()}
        po_parts = {r[0] for r in db.query(PurchaseOrder.part_no).filter(PurchaseOrder.partner_id == partner_id).all()}
        part_nos = so_parts | sh_parts | po_parts
        if part_nos:
            q = q.filter(Item.part_no.in_(part_nos))
        else:
            return []
    return q.all()


@router.get("/items/{part_no:path}", response_model=ItemOut)
def get_item(part_no: str, db: Session = Depends(get_db)):
    item = db.get(Item, part_no)
    if not item:
        raise HTTPException(404)
    return item


@router.patch("/items/{part_no:path}", response_model=ItemOut)
def update_item(part_no: str, body: ItemUpdate, db: Session = Depends(get_db)):
    item = db.get(Item, part_no)
    if not item:
        raise HTTPException(404)
    for k, v in body.model_dump(exclude_none=True).items():
        setattr(item, k, v)
    db.commit()
    db.refresh(item)
    return item


@router.post("/bulk-upload")
def bulk_upload_all(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """
    통합 엑셀 업로드 (시트 3개):
      Sheet1 '품목' : 품번, 품명, 규격, 단위, 품목구분, 표준입고가, 표준판매가, 안전재고, MOQ
      Sheet2 '거래처': 거래처코드, 거래처명, 구분(외주처/고객/공용), 사업자번호, 결제조건, 이메일
      Sheet3 '초기재고': 품번, 현재고수량, 단가(선택)
    이미 존재하면 업데이트, 없으면 신규 생성.
    """
    from datetime import date
    from app.models.ledger import StockLedger, LedgerType
    from app.services.ledger_service import get_stock

    _ITEM_TYPE_MAP = {
        '제품': 'finished', '완제품': 'finished',
        '반제품': 'semi',
        '외주품': 'outsourced', '외주': 'outsourced',
        '원자재': 'raw',
        '부자재': 'sub',
    }

    content = file.file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content))
    result = {"items": {}, "partners": {}, "stock": {}, "errors": []}

    # ── Sheet1: 품목 ──
    if '품목' in wb.sheetnames:
        ws = wb['품목']
        created, updated = 0, 0
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row[0]:
                continue
            try:
                part_no = str(row[0]).strip()
                data = {
                    "name":           str(row[1]).strip() if row[1] else "",
                    "spec":           str(row[2]).strip() if row[2] else None,
                    "unit":           str(row[3]).strip() if row[3] else "EA",
                    "item_type":      _ITEM_TYPE_MAP.get(str(row[4]).strip(), str(row[4]).strip()) if row[4] else "outsourced",
                    "std_buy_price":  float(row[5] or 0),
                    "std_sell_price": float(row[6] or 0),
                    "safety_stock":   int(row[7] or 0),
                    "moq":            int(row[8] or 1),
                }
                existing = db.get(Item, part_no)
                if existing:
                    for k, v in data.items():
                        setattr(existing, k, v)
                    updated += 1
                else:
                    db.add(Item(part_no=part_no, **data))
                    created += 1
            except Exception as e:
                result["errors"].append({"sheet": "품목", "row": row_idx, "error": str(e)})
        result["items"] = {"created": created, "updated": updated}

    # ── Sheet2: 거래처 ──
    if '거래처' in wb.sheetnames:
        ws = wb['거래처']
        created, updated = 0, 0
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row[0]:
                continue
            try:
                partner_id = str(row[0]).strip()
                data = {
                    "name":          str(row[1]).strip() if row[1] else "",
                    "partner_type":  str(row[2]).strip() if row[2] else "외주처",
                    "business_no":   str(row[3]).strip() if row[3] else None,
                    "payment_terms": str(row[4]).strip() if row[4] else None,
                    "email":         str(row[5]).strip() if row[5] else None,
                }
                existing = db.get(Partner, partner_id)
                if existing:
                    for k, v in data.items():
                        setattr(existing, k, v)
                    updated += 1
                else:
                    db.add(Partner(partner_id=partner_id, **data))
                    created += 1
            except Exception as e:
                result["errors"].append({"sheet": "거래처", "row": row_idx, "error": str(e)})
        result["partners"] = {"created": created, "updated": updated}

    db.flush()  # 품목/거래처 먼저 flush 후 재고 처리

    # ── Sheet3: 초기재고 ──
    if '초기재고' in wb.sheetnames:
        ws = wb['초기재고']
        today = date.today()
        adjusted = 0
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row[0]:
                continue
            try:
                part_no = str(row[0]).strip()
                target_qty = int(row[1] or 0)
                unit_price = float(row[2] or 0) if len(row) > 2 else 0.0

                if not db.get(Item, part_no):
                    result["errors"].append({"sheet": "초기재고", "row": row_idx, "error": f"품번 없음: {part_no}"})
                    continue

                current = get_stock(db, part_no)
                delta = target_qty - current
                if delta == 0:
                    continue

                entry = StockLedger(
                    txn_date=today,
                    part_no=part_no,
                    ledger_type=LedgerType.adjustment,
                    qty=delta,
                    unit_price=unit_price,
                    ref_type="ADJ",
                    note="초기재고 일괄 입력",
                )
                db.add(entry)
                adjusted += 1
            except Exception as e:
                result["errors"].append({"sheet": "초기재고", "row": row_idx, "error": str(e)})
        result["stock"] = {"adjusted": adjusted}

    db.commit()
    return result


@router.post("/items/import-ls")
def import_ls_items(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """
    LS Electric 납품 Excel 업로드 — 세 가지 포맷 자동 감지:
      [A] LS 납품 포맷 (7열): 품번, 품명(or Category+품명), 수량, 단가, 날짜
           → 품목 신규등록 + 수주(SO) + 출하(SH) 자동 생성. 중복 행은 스킵.
      [B] 납품현황 포맷 (26열): 헤더 1행, D=품번, F=품명, I=단가, R=단위
      [C] 단순 3열 포맷: A=품번, B=품명, C=단가
    """
    from datetime import date as date_type
    from app.models.sales import SalesOrder, Shipment, SOStatus, ShipmentStatus
    from app.services.doc_no import next_doc_no
    from app.services.ledger_service import post_ledger, get_avg_price
    from app.models.ledger import LedgerType

    content = file.file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    created, updated, skipped, errors = 0, 0, 0, []
    so_created, sh_created = 0, 0

    def _safe_float(v):
        try:
            return float(str(v).replace(',', '').strip())
        except Exception:
            return 0.0

    def _safe_date(v):
        if isinstance(v, date_type):
            return v
        try:
            from datetime import datetime
            if isinstance(v, datetime):
                return v.date()
            return date_type.fromisoformat(str(v).strip()[:10])
        except Exception:
            return date_type.today()

    def _ensure_item(part_no, name, price):
        """품목이 없으면 신규 등록, 있으면 단가만 업데이트. 생성여부 반환."""
        nonlocal created, updated
        existing = db.get(Item, part_no)
        if existing:
            if price > 0:
                existing.std_sell_price = price
            updated += 1
            return False
        else:
            db.add(Item(
                part_no=part_no, name=name, unit="EA",
                item_type="finished", std_sell_price=price,
                std_buy_price=0, safety_stock=0, moq=1,
            ))
            db.flush()
            created += 1
            return True

    first_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    h0 = str(first_row[0] or '').strip()
    h1 = str(first_row[1] or '').strip() if len(first_row) > 1 else ''
    is_ls_shipment_fmt = h0 == 'Shipment No.'

    # 포맷 A 감지: 7열짜리 + 2열이 Category 이거나 한글(품명) + Shipment No 아님
    # LS 납품 파일은 항상 6~8열이고 Shipment No. 헤더가 아님
    is_ls_delivery = (not is_ls_shipment_fmt) and (5 <= ws.max_column <= 9) and (
        h1 == 'Category' or (len(h1) >= 2 and not h1[0].isdigit())
    )

    if is_ls_delivery:
        # LS 납품 포맷: 품번(0), [Category(1),] 품명(1 or 2), 수량(2 or 3), 단가(3 or 4), ship_date(4 or 5)
        # Category 열 여부: 2행을 보고 col1이 영문이면 Category 있음
        sample = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        has_category = sample and sample[1] and str(sample[1]).replace(' ','').isascii() and len(str(sample[1])) < 20 and str(sample[1]).strip() not in ('', 'None')

        # 이미 임포트된 (part_no, ship_date, qty) 조합 체크용
        existing_sh = set()
        for row in db.query(Shipment.part_no, Shipment.ship_date, Shipment.qty).all():
            existing_sh.add((row[0], str(row[1]), row[2]))

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                skipped += 1
                continue
            part_no = str(row[0]).strip()
            if has_category:
                name  = str(row[2]).strip() if row[2] else ""
                qty   = int(_safe_float(row[3])) if row[3] else 0
                price = _safe_float(row[4]) if row[4] else 0.0
                ship_date = _safe_date(row[5] if row[5] else row[6])
            else:
                name  = str(row[1]).strip() if row[1] else ""
                qty   = int(_safe_float(row[2])) if row[2] else 0
                price = _safe_float(row[3]) if row[3] else 0.0
                ship_date = _safe_date(row[4] if row[4] else row[5])

            if not part_no or not name or qty <= 0:
                skipped += 1
                continue

            # 중복 출하 스킵
            key = (part_no, str(ship_date), qty)
            if key in existing_sh:
                skipped += 1
                continue
            existing_sh.add(key)

            try:
                _ensure_item(part_no, name, price)
                db.flush()

                # SO 생성 (closed 상태로)
                so_no = next_doc_no(db, "SO", "sales_order", "so_no")
                so = SalesOrder(
                    so_no=so_no, partner_id="LS_ELEC", part_no=part_no,
                    qty=qty, order_date=ship_date, due_date=ship_date,
                    status=SOStatus.closed,
                    note=f"LS 납품 자동등록 {ship_date}",
                )
                db.add(so)
                db.flush()
                so_created += 1

                # 출하 생성
                sh_no = next_doc_no(db, "SH", "shipment", "sh_no")
                sh = Shipment(
                    sh_no=sh_no, so_no=so_no, part_no=part_no,
                    qty=qty, ship_date=ship_date, unit_price=price,
                    status=ShipmentStatus.confirmed,
                    note=f"LS 납품 자동등록",
                )
                db.add(sh)
                db.flush()
                sh_created += 1

                # 재고 출고 반영
                avg = get_avg_price(db, part_no)
                post_ledger(
                    db, part_no, ship_date,
                    LedgerType.shipment, -qty, avg or price,
                    ref_type="SH", ref_no=sh_no,
                )
            except Exception as e:
                errors.append({"part_no": part_no, "error": str(e)})

        db.commit()
        return {
            "format": "LS납품",
            "item_created": created, "item_updated": updated,
            "so_created": so_created, "sh_created": sh_created,
            "skipped": skipped, "errors": errors,
        }

    elif is_ls_shipment_fmt:
        # 납품현황 포맷: D(3)=품번, F(5)=품명, I(8)=단가, R(17)=단위
        items_map: dict[str, dict] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            part_no = str(row[3]).strip() if row[3] else ""
            name    = str(row[5]).strip() if row[5] else ""
            price   = _safe_float(row[8]) if len(row) > 8 else 0.0
            unit    = str(row[17]).strip().upper() if len(row) > 17 and row[17] else "EA"
            if not part_no or not name:
                skipped += 1
                continue
            if unit not in ('EA', 'KG', 'M', 'SET', 'L'):
                unit = 'EA'
            if part_no not in items_map:
                items_map[part_no] = {'name': name, 'price': price, 'unit': unit}
            else:
                if price > 0:
                    items_map[part_no]['price'] = price

        for part_no, d in items_map.items():
            try:
                existing = db.get(Item, part_no)
                if existing:
                    existing.name = d['name']
                    if d['price'] > 0:
                        existing.std_buy_price = d['price']
                    updated += 1
                else:
                    db.add(Item(
                        part_no=part_no, name=d['name'], unit=d['unit'],
                        item_type="outsourced", std_buy_price=d['price'],
                        std_sell_price=0, safety_stock=0, moq=1,
                    ))
                    created += 1
            except Exception as e:
                errors.append({"part_no": part_no, "error": str(e)})
    else:
        # 단순 3열 포맷: A=품번, B=품명, C=단가
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
            if not row[0] and not row[1]:
                skipped += 1
                continue
            part_no = str(row[0]).strip() if row[0] else ""
            name    = str(row[1]).strip() if row[1] else ""
            if not part_no or part_no in ('품번', 'Part No', 'PART_NO'):
                skipped += 1
                continue
            if not name:
                skipped += 1
                continue
            price = _safe_float(row[2]) if len(row) > 2 and row[2] else 0.0
            try:
                existing = db.get(Item, part_no)
                if existing:
                    existing.name = name
                    if price > 0:
                        existing.std_buy_price = price
                    updated += 1
                else:
                    db.add(Item(
                        part_no=part_no, name=name, unit="EA",
                        item_type="outsourced", std_buy_price=price,
                        std_sell_price=0, safety_stock=0, moq=1,
                    ))
                    created += 1
            except Exception as e:
                errors.append({"row": row_idx, "part_no": part_no, "error": str(e)})

    db.commit()
    return {"created": created, "updated": updated, "skipped": skipped, "errors": errors}


@router.post("/items/bulk-upload")
def bulk_upload_items(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """하위호환: 품목만 단독 업로드 (열: 품번, 품명, 규격, 단위, 품목구분, 표준입고가, 표준판매가, 안전재고, MOQ)"""
    wb = openpyxl.load_workbook(io.BytesIO(file.file.read()))
    ws = wb.active
    created, updated, errors = 0, 0, []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:
            continue
        try:
            part_no = str(row[0]).strip()
            data = {
                "name":           str(row[1]).strip() if row[1] else "",
                "spec":           str(row[2]).strip() if row[2] else None,
                "unit":           str(row[3]).strip() if row[3] else "EA",
                "item_type":      str(row[4]).strip() if row[4] else "outsourced",
                "std_buy_price":  float(row[5] or 0),
                "std_sell_price": float(row[6] or 0),
                "safety_stock":   int(row[7] or 0),
                "moq":            int(row[8] or 1),
            }
            existing = db.get(Item, part_no)
            if existing:
                for k, v in data.items():
                    setattr(existing, k, v)
                updated += 1
            else:
                db.add(Item(part_no=part_no, **data))
                created += 1
        except Exception as e:
            errors.append({"row": row_idx, "error": str(e)})
    db.commit()
    return {"created": created, "updated": updated, "errors": errors}




@router.post("/partners", response_model=PartnerOut)
def create_partner(body: PartnerCreate, db: Session = Depends(get_db)):
    if db.get(Partner, body.partner_id):
        raise HTTPException(409, f"이미 존재: {body.partner_id}")
    partner = Partner(**body.model_dump())
    db.add(partner)
    db.commit()
    db.refresh(partner)
    return partner


@router.get("/partners", response_model=list[PartnerOut])
def list_partners(db: Session = Depends(get_db)):
    return db.query(Partner).filter(Partner.active == True).all()


@router.get("/partners/{partner_id}", response_model=PartnerOut)
def get_partner(partner_id: str, db: Session = Depends(get_db)):
    p = db.get(Partner, partner_id)
    if not p:
        raise HTTPException(404)
    return p


@router.patch("/partners/{partner_id}", response_model=PartnerOut)
def update_partner(partner_id: str, body: PartnerUpdate, db: Session = Depends(get_db)):
    p = db.get(Partner, partner_id)
    if not p:
        raise HTTPException(404)
    for k, v in body.model_dump(exclude_none=True).items():
        setattr(p, k, v)
    db.commit()
    db.refresh(p)
    return p
